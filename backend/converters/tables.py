import csv
import io

import pymupdf

MAX_TABLES = 200

# The text-detection strategy will happily find a "table" in any text that
# happens to line up in columns — reference lists and figure labels both
# qualify. These thresholds reject those without touching real tables.
MAX_COLUMNS = 12
MIN_FILL_RATIO = 0.5
MAX_SPLIT_WORD_RATIO = 0.25


def _fill_ratio(rows: list[list[str]]) -> float:
    cells = sum(len(row) for row in rows)
    if not cells:
        return 0.0
    return sum(1 for row in rows for cell in row if cell) / cells


def _splits_words(rows: list[list[str]]) -> bool:
    """
    True when cell boundaries fall inside words.

    Column detection applied to prose slices continuous text mid-token, giving
    neighbouring cells like "Kyunghyun Cho, Bar" and "t van Merrienb". A real
    table's cells are self-contained, so this only fires on misread text.
    """
    splits = 0
    pairs = 0
    for row in rows:
        for left, right in zip(row, row[1:]):
            if not left or not right:
                continue
            pairs += 1
            if left[-1].isalnum() and right[0].islower():
                splits += 1
    return pairs > 0 and splits / pairs > MAX_SPLIT_WORD_RATIO


def _is_plausible_table(rows: list[list[str]], from_lines: bool) -> bool:
    """
    Shape checks apply to every candidate, however it was detected.

    Ruled detection isn't automatically trustworthy: a figure drawn with vector
    grid lines — an attention heatmap, say — reads as a 50-column table of
    single characters. Only the split-word test is limited to text detection,
    since ruled cells don't slice words.
    """
    if len(rows) < 2:
        return False
    if max(len(row) for row in rows) > MAX_COLUMNS:
        return False
    if _fill_ratio(rows) < MIN_FILL_RATIO:
        return False
    return from_lines or not _splits_words(rows)


def _clean_rows(rows: list[list]) -> list[list[str]]:
    """Normalise cells to strings and drop rows that are entirely empty.

    The text-based detection strategy interleaves blank rows between real ones,
    so this keeps output usable regardless of which strategy found the table.
    """
    cleaned = []
    for row in rows:
        cells = [("" if cell is None else str(cell)).strip() for cell in row]
        if any(cells):
            cleaned.append(cells)
    return cleaned


def find_tables(file_bytes: bytes) -> list[tuple[int, list[list[str]]]]:
    """Return (page_number, rows) for every table found in the PDF."""
    document = pymupdf.open(stream=file_bytes, filetype="pdf")
    try:
        found: list[tuple[int, list[list[str]]]] = []

        for number, page in enumerate(document, start=1):
            # Ruled tables detect cleanly from their lines; borderless ones
            # only show up under the text strategy, which is far noisier — so
            # anything it returns has to survive the plausibility checks.
            ruled = page.find_tables(strategy="lines").tables
            candidates = [(table, True) for table in ruled]
            if not ruled:
                candidates = [
                    (table, False) for table in page.find_tables(strategy="text").tables
                ]

            for table, from_lines in candidates:
                rows = _clean_rows(table.extract())
                if _is_plausible_table(rows, from_lines):
                    found.append((number, rows))
                if len(found) >= MAX_TABLES:
                    return found

        return found
    finally:
        document.close()


def tables_to_csv(rows: list[list[str]]) -> str:
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(rows)
    return buffer.getvalue()


def tables_to_xlsx(tables: list[tuple[int, list[list[str]]]]) -> bytes:
    """One worksheet per table, named by the page it came from."""
    from openpyxl import Workbook  # imported lazily to keep idle memory down

    workbook = Workbook()
    workbook.remove(workbook.active)

    for index, (page_number, rows) in enumerate(tables, start=1):
        sheet = workbook.create_sheet(title=f"p{page_number}-table{index}"[:31])
        for row in rows:
            sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
